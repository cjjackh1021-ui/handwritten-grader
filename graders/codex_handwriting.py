"""Agent-native handwriting workflow. All grading comes from the current chat, never an API."""
from __future__ import annotations

import argparse
import copy
from dataclasses import asdict
import hashlib
import json
import math
from pathlib import Path
import sys
import time
import zipfile

from . import annotations, extract, local_workers, review, stamp, summary


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def write(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    temp.replace(path)
    return str(path.resolve())


def fresh_output(path):
    path = Path(path).resolve()
    if path.exists() and any(path.iterdir()):
        raise ValueError("输出目录必须为空；请为每次修订使用新目录，保留已有结果和人工修改表")
    path.mkdir(parents=True, exist_ok=True)
    return path


def fingerprint(q):
    keys = ("id", "question", "student_answer", "answer_state", "full", "reference", "rubric", "regions")
    return hashlib.sha256(json.dumps({k: q.get(k) for k in keys}, ensure_ascii=False, sort_keys=True).encode()).hexdigest()


def verify_sources(reports):
    for r in reports:
        for source in r.get("inputs", []):
            path = Path(source["path"])
            if not path.is_file() or extract.sha256_file(path) != source["sha256"]:
                raise ValueError(f"原稿已移动或变更，请重新提取：{path}")


def prepare(manifest_path, output, ocr_mode="hybrid"):
    manifest_path = Path(manifest_path).resolve()
    manifest = read(manifest_path)
    submissions = manifest.get("submissions", [])
    if not submissions:
        raise ValueError("manifest.submissions 不能为空")
    resolved, used = [], set()
    for row in submissions:
        files = []
        for name in row.get("files", []):
            path = Path(name)
            path = (manifest_path.parent / path).resolve() if not path.is_absolute() else path.resolve()
            if not path.is_file() or path.suffix.lower() not in extract.IMAGE_EXT | {".pdf"}:
                raise ValueError(f"作答须为已存在的图片或 PDF：{path}")
            if path in used:
                raise ValueError(f"同一文件不能归入两份作答：{path}")
            used.add(path)
            files.append(path)
        if not files:
            raise ValueError("每份作答至少包含一个文件；files 顺序即页序")
        resolved.append((row, files))
    out = fresh_output(output)
    reports = []
    for row, files in resolved:
        inputs = [{"path": str(p), "sha256": extract.sha256_file(p)} for p in files]
        rid = annotations.stable_id(*[f"{v['path']}:{v['sha256']}" for v in inputs])
        entry = dict(report_id=rid, inputs=inputs, student_name=row.get("student_name", files[0].stem),
                     student_id=str(row.get("student_id", "")), source_path=str(files[0]), source_file=files[0].name)
        try:
            parts = []
            for path in files:
                print(f"正在提取 {path.name}", file=sys.stderr, flush=True)
                inspect = extract.inspect_pdf if path.suffix.lower() == ".pdf" else extract.inspect_image
                parts.append(inspect(path, out / "_work", ocr_enabled=ocr_mode != "off", ocr_mode=ocr_mode))
            report = extract.merge_reports(parts, files[0].parent)
            if not report.pages or any(not p.render_path for p in report.pages):
                raise ValueError("页面渲染不完整，不能确认定位；请检查原稿")
            report.student_name = entry["student_name"]
            report.student_id = entry["student_id"]
            entry.update(status="success", report_snapshot=asdict(report), pages=[asdict(p) for p in report.pages],
                         evidence=annotations.evidence(report), source_sha256=report.source_sha256)
        except Exception as exc:
            entry.update(status="error", error=str(exc), pages=[])
        reports.append(entry)
        write(out / "packet.json", dict(schema_version=2, workflow="codex-handwriting", ocr_mode=ocr_mode,
                                        reference_notes=manifest.get("reference_notes", ""), reports=reports))
    return {"packet": str(out / "packet.json"), "reports": len(reports),
            "failures": [{"report_id": r["report_id"], "error": r["error"]} for r in reports if r["status"] == "error"]}


def grade_template(reports, only=None):
    return {"grades": [dict(report_id=r["report_id"], question_id=q["id"], question_digest=fingerprint(q),
                            score=None, grading_confirmed=False, comment="", reason="")
                       for r in reports for q in r.get("details", {}).get("questions", [])
                       if only is None or (r["report_id"], q["id"]) in only]}


def locate(packet_path, questions_path, output):
    packet = read(packet_path)
    reports = copy.deepcopy(packet["reports"])
    verify_sources(reports)
    rows = read(questions_path).get("reports", [])
    by_id = {r["report_id"]: r for r in rows}
    expected = {r["report_id"] for r in reports if r["status"] == "success"}
    if len(by_id) != len(rows) or set(by_id) != expected:
        raise ValueError("questions.reports 须一一对应所有提取成功的作答 ID")
    for entry in reports:
        if entry["status"] != "success":
            continue
        report = extract.report_from_dict(entry["report_snapshot"])
        raw_questions = by_id[entry["report_id"]].get("questions", [])
        if not raw_questions:
            raise ValueError("每份成功作答至少需一个小题；未识别题目请保留为待复核")
        questions, numbers = [], set()
        for raw in raw_questions:
            no = str(raw["no"]).strip()
            full = float(raw["full"])
            if not no or no in numbers or not math.isfinite(full) or full <= 0:
                raise ValueError("小题编号须非空且唯一；满分须为正有限数")
            numbers.add(no)
            if raw.get("answer_state") not in ("present", "blank", "unreadable"):
                raise ValueError("answer_state 须明确为 present、blank 或 unreadable")
            valid_pages = {p.page_num for p in report.pages}
            for region in raw.get("regions", []):
                if region.get("page") not in valid_pages or region.get("role") not in ("answer", "question") or not annotations.valid_box(region.get("bbox")):
                    raise ValueError("题目区域页码、角色或坐标无效")
            blocks = {b["id"] for p in entry["evidence"] for b in p["blocks"]}
            if any(b not in blocks for role in ("answer", "question") for b in raw.get(role + "_block_ids", [])):
                raise ValueError("题目关联了不存在的 OCR 块 ID")
            q = dict(no=no, question=str(raw.get("question", "")), student_answer=str(raw.get("student_answer", "")),
                     full=full, type=raw.get("type", "short"), reference=str(raw.get("reference", "")),
                     rubric=raw.get("rubric", ""), reference_source=raw.get("reference_source", "unspecified"))
            annotations.prepare_question(q, raw, report, visual=True)
            q["id"] = annotations.stable_id(entry["report_id"], no)
            q.update(score=0.0, grading_review=True, score_pending=True, comment="待评分", reason="尚未评分")
            if not q["question"].strip():
                q["recognition_review"] = True
                q["review_reasons"].append("题干尚未确认")
            q["review_reasons"].extend(str(s) for s in raw.get("review_reasons", []))
            if raw.get("recognition_review") is True:
                q["recognition_review"] = True
            if raw.get("position_review") is True:
                q["position_review"] = True
            questions.append(q)
        annotations.check_region_collisions(questions)
        for q in questions:
            cropped = annotations.crop_report(report, q)
            q["crop_paths"] = [p.render_path for p in cropped.pages]
        entry.update(task_type="exam", details={"questions": questions}, graded_by="codex-current-session")
        review.aggregate(entry)
    out = fresh_output(output)
    return {"layout": write(out / "layout.json", {**packet, "reports": reports}),
            "grading_template": write(out / "grades.template.json", grade_template(reports))}


def apply_grades(reports, grades, require_all=True):
    reports = copy.deepcopy(reports)
    pairs = {(r["report_id"], q["id"]): q for r in reports for q in r.get("details", {}).get("questions", [])}
    seen = set()
    for row in grades.get("grades", []):
        key = (row["report_id"], row["question_id"])
        if key not in pairs or key in seen:
            raise ValueError("评分包含未知或重复的作答/小题 ID")
        seen.add(key)
        q = pairs[key]
        if row.get("question_digest") != fingerprint(q):
            raise ValueError("题干、转写、分值或区域已变化，请针对最新题目重新评分")
        confirmed = row.get("grading_confirmed") is True
        if confirmed and (row.get("score") is None or not str(row.get("reason", "")).strip()):
            raise ValueError("确认评分须有分数和给分依据")
        value = review.valid_score(row["score"], q["full"]) if row.get("score") is not None else 0.0
        if q["answer_state"] == "blank" and not q["recognition_review"] and value != 0:
            raise ValueError("确认空白的题目得分须为 0")
        # An unconfirmed/illegible answer contributes no asserted score, and never gets a cross.
        pending = not confirmed or q["recognition_review"]
        q.update(score=0.0 if pending else value, score_pending=pending, grading_review=pending,
                 comment=str(row.get("comment", "")), reason=str(row.get("reason", "")), graded_by="codex-current-session")
        q.pop("warning", None)
        q.pop("error", None)
    if require_all and seen != set(pairs):
        raise ValueError("首次评分必须包含所有小题，疑题使用 score=null、grading_confirmed=false")
    for r in reports:
        if any(r["report_id"] == rid for rid, _ in seen):
            r.pop("manual_total_adjustment", None)
        review.aggregate(r)
    return reports


def export(reports, output, style=None):
    verify_sources(reports)
    out = fresh_output(output)
    reports = copy.deepcopy(reports)
    good = [r for r in reports if r.get("status") == "success"]
    for r in good:
        r["handwriting_style"] = style or r.get("handwriting_style", "liguofu")
        r["annotate_questions"] = True
        review.aggregate(r)
        for q in r.get("details", {}).get("questions", []):
            q["score_pending"] = bool(q.get("recognition_review") or q.get("grading_review"))
    result = stamp.stamp_batch(good, out)
    # Include extraction failures in the persisted batch, never silently omit a student.
    write(out / "batch_results.json", {"schema_version": 2, "workflow": "codex-handwriting", "reports": reports})
    xlsx = summary.build_workbook(reports, out / "批改结果.xlsx")
    from openpyxl import load_workbook
    workbook = load_workbook(xlsx)
    notes = workbook["编辑说明"]
    notes.cell(2, 1, "转写、区域和标记可在 Codex 使用本技能 review 命令复核；文字修改后只重批该题。")
    notes.cell(4, 1, "未确认题目的 0 是暂计占位，不能当作确认零分；待复核题不会落笔。")
    workbook.save(xlsx)
    workbook.close()
    pending = []
    failures = []
    for r in reports:
        if r.get("error") or r.get("stamp_error"):
            failures.append({"report_id": r["report_id"], "error": r.get("error") or r["stamp_error"]})
        for q in r.get("details", {}).get("questions", []):
            if q.get("review_required"):
                pending.append(dict(report_id=r["report_id"], student_name=r.get("student_name"), question=q,
                                    evidence=r.get("evidence", [])))
    write(out / "review_pending.json", {"questions": pending, "failures": failures})
    write(out / "grades.template.json", grade_template(reports, {(p["report_id"], p["question"]["id"]) for p in pending}))
    # Rewrite once so Excel, batch and PNG previews describe the same post-placement state.
    with zipfile.ZipFile(result["zip_path"], "w", zipfile.ZIP_DEFLATED) as archive:
        for r in good:
            for path in [r.get("marked_pdf_path", ""), *r.get("marked_pngs", [])]:
                if path and Path(path).is_file():
                    archive.write(path, Path(path).relative_to(out))
        for name in ("batch_results.json", "批改结果.xlsx", "review_pending.json"):
            archive.write(out / name, name)
    return {**result, "excel": xlsx, "review": str(out / "review_pending.json"),
            "pending_questions": len(pending), "failures": failures,
            "marked": sum(bool(r.get("marked_pdf_path")) for r in good)}


def revise(reports, changes):
    reports = copy.deepcopy(reports)
    lookup = {r["report_id"]: i for i, r in enumerate(reports)}
    for change in changes.get("changes", []):
        rid, qid = change["report_id"], change["question_id"]
        if rid not in lookup:
            raise ValueError("作答 ID 不存在")
        idx = lookup[rid]
        entry = reports[idx]
        q = next((q for q in entry.get("details", {}).get("questions", []) if q["id"] == qid), None)
        if q is None:
            raise ValueError("小题 ID 不存在")
        if change.get("action") == "transcript":
            state, text = change.get("answer_state"), str(change.get("student_answer", "")).strip()
            if state not in ("present", "blank", "unreadable") or (state == "blank" and text) or (state == "present" and not text):
                raise ValueError("转写与作答状态不一致")
            entry.setdefault("review_history", []).append(dict(time=time.time(), question_id=qid, action="transcript",
                                                                before=copy.deepcopy(q)))
            q.update(student_answer=text, answer_state=state, recognition_review=state == "unreadable" or not q.get("question", "").strip(),
                     grading_review=True, score_pending=True, score=0.0, comment="转写已修正，待重批", reason="")
            q.pop("warning", None)
            q.pop("error", None)
            entry.pop("manual_total_adjustment", None)
            q["review_reasons"] = ["转写已修正，请用当前 Codex 模型只重批该题"]
            review.aggregate(entry)
        else:
            reports[idx] = review.edit_question(entry, qid, change)
            q = next(q for q in reports[idx]["details"]["questions"] if q["id"] == qid)
            q["score_pending"] = bool(q.get("recognition_review") or q.get("grading_review"))
            if change.get("action") == "location":
                report = extract.report_from_dict(reports[idx]["report_snapshot"])
                q["crop_paths"] = [p.render_path for p in annotations.crop_report(report, q).pages]
    return reports


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    commands.add_parser("doctor")
    p = commands.add_parser("prepare")
    p.add_argument("--manifest", type=Path, required=True)
    p.add_argument("--output", type=Path, required=True)
    p.add_argument("--ocr-mode", choices=("hybrid", "paddle", "tesseract", "off"), default="hybrid")
    p = commands.add_parser("locate")
    p.add_argument("--packet", type=Path, required=True)
    p.add_argument("--questions", type=Path, required=True)
    p.add_argument("--output", type=Path, required=True)
    for name in ("export", "apply-grades", "review", "regenerate"):
        p = commands.add_parser(name)
        p.add_argument("--batch", type=Path, required=True, help="layout.json or batch_results.json")
        p.add_argument("--output", type=Path, required=True)
        p.add_argument("--style", choices=("liguofu", "kaishu"), default=None)
        argument = "grades" if name in ("export", "apply-grades") else "changes" if name == "review" else "xlsx"
        p.add_argument("--" + argument, type=Path, required=True)
    args = parser.parse_args(argv)
    try:
        if args.command == "doctor":
            result = dict(project=str(local_workers.ROOT), python=sys.executable, components=local_workers.capabilities(),
                          grading="current Codex session; no external grading API", model_downloads=False)
        elif args.command == "prepare":
            result = prepare(args.manifest, args.output, args.ocr_mode)
        elif args.command == "locate":
            result = locate(args.packet, args.questions, args.output)
        else:
            reports = read(args.batch)["reports"]
            verify_sources(reports)
            if args.command in ("export", "apply-grades"):
                reports = apply_grades(reports, read(args.grades), require_all=args.command == "export")
            elif args.command == "review":
                reports = revise(reports, read(args.changes))
            else:
                reports = summary.merge_edits_into_batch(summary.read_workbook(args.xlsx), args.batch)
            result = export(reports, args.output, args.style)
        print(json.dumps(result, ensure_ascii=False, allow_nan=False), flush=True)
        return 1 if result.get("failures") else 0
    except (ValueError, KeyError, TypeError, OSError) as exc:
        print(json.dumps({"error": str(exc)}, ensure_ascii=False), file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
