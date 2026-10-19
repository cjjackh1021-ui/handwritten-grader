"""DOCX -> PDF 转换。

优先用 LibreOffice/soffice 生成高保真 staging PDF；若环境无 LibreOffice，
降级为用 python-docx 仅提取文本 (供纯文本评分, 无法盖分)。

移植自 skill 的 convert_word_reports.py，保留来源哈希作为溯源。
"""

from __future__ import annotations

import hashlib
import subprocess
from dataclasses import dataclass
from pathlib import Path

from . import config


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


@dataclass
class ConversionResult:
    source_path: str
    source_sha256: str
    source_file: str
    pdf_path: str | None  # 转换出的 staging PDF；降级时为 None
    text: str | None      # 降级模式下提取到的纯文本
    error: str | None = None


def _libreoffice_convert(soffice: str, source: Path, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    target = out_dir / f"{source.stem}.pdf"
    completed = subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(source)],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=180,
    )
    if completed.returncode or not target.is_file():
        raise RuntimeError(
            completed.stderr.strip() or completed.stdout.strip() or "Word 转 PDF 失败"
        )
    return target


def _docx_text(source: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("未安装 python-docx，且无 LibreOffice，无法处理 DOCX") from exc
    document = docx.Document(str(source))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def convert_docx(source: Path, staging_dir: Path) -> ConversionResult:
    """转换单个 DOCX。"""
    source = source.resolve()
    soffice = config.find_soffice()
    base = ConversionResult(
        source_path=str(source),
        source_sha256=sha256(source),
        source_file=source.name,
        pdf_path=None,
        text=None,
    )
    if soffice:
        try:
            pdf = _libreoffice_convert(soffice, source, staging_dir)
            base.pdf_path = str(pdf.resolve())
            return base
        except Exception as exc:  # noqa: BLE001 - 记录后尝试降级
            base.error = f"LibreOffice 转换失败: {exc}"
    # 降级：仅取文本
    try:
        base.text = _docx_text(source)
        if base.error:
            base.error += "；已降级为纯文本提取"
    except Exception as exc:  # noqa: BLE001
        base.error = (base.error + "；" if base.error else "") + str(exc)
    return base
