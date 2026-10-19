#!/usr/bin/env python3
"""学生分组 Excel 汇总 (openpyxl, 替代原 build_summary.mjs, 去除 OAI 专有依赖)。

生成:
    python build_summary.py --input output/batch_results.json --output output/批改结果.xlsx

改表回读 (重生成评分 JSON):
    python build_summary.py --from-xlsx output/批改结果.xlsx \
        --batch output/batch_results.json --emit output/grades.json
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ["姓名", "学号", "分数", "分维度得分", "评语", "给分依据", "抄袭嫌疑", "问题页码", "问题原因", "来源文件"]
WIDTHS = [14, 16, 8, 34, 30, 40, 10, 12, 24, 28]


def _dimension_summary(item):
    if item.get("dimension_summary"):
        return item["dimension_summary"]
    parts = []
    for d in item.get("dimensions") or []:
        name, score, full = d.get("name", ""), d.get("score"), d.get("full")
        if score is None:
            continue
        s = f"{name}{score}"
        if full is not None:
            s += f"/{full}"
        parts.append(s)
    return "; ".join(parts)


def _sort_key(item):
    return (str(item.get("student_id") or "~"), str(item.get("student_name", "")), str(item.get("source_file", "")))


def build_workbook(reports, out_xlsx: Path):
    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    rows = sorted([r for r in reports if r.get("status", "success") == "success"], key=_sort_key)
    if not rows:
        raise SystemExit("没有可写入 Excel 的成功报告")
    wb = Workbook()
    ws = wb.active
    ws.title = "批改结果"
    ws.sheet_view.showGridLines = False
    ws.append(HEADERS)
    hfill, hfont = PatternFill("solid", fgColor="667EEA"), Font(color="FFFFFF", bold=True)
    for col, w in enumerate(WIDTHS, 1):
        cell = ws.cell(1, col)
        cell.fill, cell.font = hfill, hfont
        cell.alignment = Alignment(vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = w
    prev, band = None, False
    for item in rows:
        plag = item.get("plagiarism_info", {}) or {}
        key = (item.get("student_id"), item.get("student_name"))
        if key != prev:
            band = not band
            prev = key
        ws.append([item.get("student_name", ""), item.get("student_id", ""), int(item.get("score", 0)),
                   _dimension_summary(item), item.get("comment", ""), item.get("reason", ""),
                   "是" if plag.get("plagiarism") else "否",
                   item.get("problem_pages", "无"), item.get("problem_reasons", "无"), item.get("source_file", "")])
        r = ws.max_row
        fill = PatternFill("solid", fgColor="F3F5FF") if band else None
        for col in range(1, len(HEADERS) + 1):
            cell = ws.cell(r, col)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fill:
                cell.fill = fill
        ws.cell(r, 2).number_format = "@"
    ws.freeze_panes = "A2"
    wb.save(str(out_xlsx))
    return str(out_xlsx.resolve())


def read_workbook(xlsx_path: Path):
    wb = load_workbook(str(xlsx_path))
    ws = wb.active
    header = [c.value for c in ws[1]]
    idx = {n: header.index(n) for n in HEADERS if n in header}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or all(v in (None, "") for v in row):
            continue
        g = lambda n, d="": (row[idx[n]] if n in idx and idx[n] < len(row) and row[idx[n]] is not None else d)
        out.append({"student_name": str(g("姓名")), "student_id": str(g("学号")),
                    "score": int(float(g("分数", 0) or 0)), "comment": str(g("评语")),
                    "reason": str(g("给分依据")), "source_file": str(g("来源文件"))})
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", type=Path, help="batch_results.json -> xlsx")
    ap.add_argument("--output", type=Path, help="输出 xlsx 路径")
    ap.add_argument("--from-xlsx", dest="from_xlsx", type=Path, help="已编辑 xlsx")
    ap.add_argument("--batch", type=Path, help="原 batch_results.json (回读时匹配用)")
    ap.add_argument("--emit", type=Path, help="回读产出的 grades.json")
    args = ap.parse_args()

    if args.from_xlsx:
        rows = read_workbook(args.from_xlsx)
        by_file = {}
        if args.batch and args.batch.is_file():
            by_file = {r.get("source_file"): r for r in json.loads(args.batch.read_text(encoding="utf-8")).get("reports", [])}
        merged = []
        for edit in rows:
            base = dict(by_file.get(edit["source_file"], {}))
            base.update({**edit, "status": "success"})
            merged.append(base)
        emit = args.emit or Path("output/grades.json")
        emit.parent.mkdir(parents=True, exist_ok=True)
        emit.write_text(json.dumps({"reports": merged}, ensure_ascii=False, indent=2), encoding="utf-8")
        print(json.dumps({"emitted": str(emit.resolve()), "reports": len(merged)}, ensure_ascii=False))
        return 0

    if not args.input or not args.output:
        ap.error("生成模式需要 --input 与 --output")
    reports = json.loads(args.input.read_text(encoding="utf-8")).get("reports", [])
    path = build_workbook(reports, args.output)
    print(json.dumps({"workbook": path, "rows": len(reports)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
