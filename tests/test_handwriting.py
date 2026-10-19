"""Regression tests for real image uploads, review, geometry and handwritten export."""
import copy
from dataclasses import asdict
import hashlib
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from PIL import Image, ImageDraw, ImageFont, ImageChops
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, ArrayObject, FloatObject, RectangleObject
from reportlab.pdfgen import canvas
from graders import annotations, config, extract, handwriting, local_workers, ocr, review, stamp, summary


class HandwritingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.path = self.root / "student.png"
        image = Image.new("RGB", (800, 1000), "white")
        draw = ImageDraw.Draw(image)
        draw.text((70, 230), "1. 2+3=5", fill="black", font=ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 35))
        image.save(self.path)
        self.report = extract.inspect_image(self.path, self.root / "_work", ocr_enabled=False)
        self.entry = {"source_path": str(self.path), "source_file": "student.png", "report_id": "student1",
                      "report_snapshot": asdict(self.report), "pages": [asdict(p) for p in self.report.pages],
                      "student_name": "学生", "student_id": "001", "status": "success", "score": 3.0, "full_score": 3.0,
                      "comment": "全部正确", "details": {"questions": [{"id": "q1", "no": "1", "type": "fill", "score": 3.0, "full": 3.0,
                      "student_answer": "5", "answer_state": "present", "pages": [1],
                      "regions": [{"page": 1, "role": "answer", "bbox": [.08,.22,.6,.31]}]}]}}

    def test_mark_decisions_fractional_and_uncertain(self):
        q = dict(score=1.5, full=2)
        self.assertEqual(annotations.mark_for(q), "partial")
        self.assertEqual(annotations.mark_for(dict(score=2, full=2)), "correct")
        self.assertEqual(annotations.mark_for(dict(score=0, full=2, answer_state="blank")), "incorrect")
        for flag in ("recognition_review", "position_review", "grading_review", "error", "warning"):
            self.assertIsNone(annotations.mark_for({**q, flag: True}))
        for value in (float("nan"), float("inf"), -1, 3, "invalid"):
            self.assertIsNone(annotations.mark_for(dict(score=value, full=2)))

    def test_unknown_empty_is_not_confirmed_blank(self):
        q = dict(no="1", student_answer="", pages=[1])
        annotations.prepare_question(q, {}, self.report, visual=False)
        self.assertEqual(q["answer_state"], "unreadable")
        self.assertTrue(q["recognition_review"])
        annotations.prepare_question(q, dict(answer_state="blank",transcription_confirmed=True),self.report,visual=True)
        self.assertFalse(q["recognition_review"])

    def test_conflicting_blank_state_is_pending(self):
        q = dict(no="1",student_answer="5",pages=[1])
        annotations.prepare_question(q, dict(answer_state="blank",transcription_confirmed=True),self.report,visual=True)
        self.assertTrue(q["recognition_review"])

    def test_native_pdf_transcription_does_not_require_vision(self):
        self.report.pages[0].needs_visual = False
        q = dict(no="1",student_answer="5",pages=[1])
        annotations.prepare_question(q, {}, self.report)
        self.assertFalse(q["recognition_review"])

    def test_ocr_missing_retains_good_engine_and_no_fake_confidence(self):
        good = dict(engine="chandra",text="5",blocks=[dict(id="c1",text="5",bbox=[.1,.1,.2,.2],confidence=None)])
        def worker(engine, *_):
            if engine == "paddle": raise local_workers.WorkerError("not ready")
            return copy.deepcopy(good)
        with patch.object(local_workers, "request", side_effect=worker):
            result = ocr.recognize(self.path, self.root, "hybrid")
        self.assertFalse(result["enhancement_complete"])
        self.assertEqual(result["text"], "5")
        self.assertIsNone(result["blocks"][0]["confidence"])

    def test_ocr_success_is_content_cached(self):
        result = dict(engine="paddle",text="5",blocks=[dict(id="p0",text="5",bbox=[.1,.1,.2,.2],confidence=.99)])
        with patch.object(local_workers,"request",return_value=result) as call:
            ocr.recognize(self.path,self.root,"paddle")
            second = ocr.recognize(self.path,self.root,"paddle")
        self.assertEqual(call.call_count,1)
        self.assertTrue(second["results"][0]["cached"])

    def test_invalid_ocr_boxes_are_not_cached_as_success(self):
        result=dict(engine="paddle",text="5",blocks=[dict(id="p0",text="5",bbox=[-.1,0,2,1],confidence=.99)])
        with patch.object(local_workers,"request",return_value=result),patch.object(ocr,"tesseract",side_effect=RuntimeError("missing")):
            recognized=ocr.recognize(self.path,self.root,"paddle")
        self.assertFalse(recognized["enhancement_complete"])
        self.assertEqual(recognized["results"],[])

    def test_export_failure_retains_scores_and_questions(self):
        with patch.object(handwriting,"render",side_effect=local_workers.WorkerError("font missing")):
            result=stamp.stamp_batch([self.entry],self.root)
        saved=json.loads(Path(result["batch_results_path"]).read_text(encoding="utf-8"))["reports"][0]
        self.assertEqual(saved["score"],3)
        self.assertEqual(saved["details"]["questions"][0]["id"],"q1")
        self.assertEqual(saved["marked_pdf_path"],"")
        self.assertIn("font missing",saved["stamp_error"])

    def test_overlapping_question_regions_require_position_review(self):
        first=copy.deepcopy(self.entry["details"]["questions"][0]);second=copy.deepcopy(first);second["id"]="q2"
        annotations.check_region_collisions([first,second])
        self.assertTrue(first["position_review"])
        self.assertTrue(second["position_review"])
        entry={"details":{"questions":[first,second]}}
        annotations.summarize(entry)
        self.assertTrue(entry["review_required"])
        self.assertFalse(entry["score_provisional"])

    def test_benchmark_metrics_use_student_transcription_not_correct_answer(self):
        from eval.handwriting_eval import score_page
        label=dict(text="2+3=6",questions=[dict(no="1",answer_bbox=[.1,.1,.5,.5])],formulas=[dict(text="2+3=6",bbox=[.1,.1,.5,.5])])
        prediction=dict(text="2+3=5",formula_predictions=["2+3=5"],questions=[dict(no="1",regions=[dict(role="answer",bbox=[.1,.1,.5,.5])])])
        result=score_page(label,prediction)
        self.assertEqual(result["char_errors"],1)
        self.assertEqual(result["formula_correct"],0)
        self.assertEqual(result["location_correct"],1)

    def test_model_score_cannot_clear_recognition_issue(self):
        q=self.entry["details"]["questions"][0];q.update(recognition_review=True,answer_state="unreadable")
        result=review.edit_question(self.entry,"q1",dict(action="score",score=2))
        self.assertTrue(result["score_provisional"])
        self.assertIsNone(result["details"]["questions"][0]["mark"])

    def test_worker_timeout_stops_process_and_reports_reason(self):
        import threading
        released=threading.Event()
        class Output:
            def readline(self):released.wait(1);return ""
            def close(self):pass
        class Process:
            stdin=io.StringIO();stdout=Output()
            def poll(self):return None
            def terminate(self):released.set()
            def wait(self,timeout=5):return 0
        worker=local_workers.Worker("paddle")
        self.addCleanup(worker.stop)
        with patch.object(local_workers,"capabilities",return_value={"paddle":{"available":True}}),patch.object(local_workers.subprocess,"Popen",return_value=Process()):
            with self.assertRaises(local_workers.WorkerError):worker.request({},timeout=.01)
        self.assertIsNone(worker.process)
        self.assertTrue(released.is_set())
        self.assertIn("未完成",worker.last_error)

    def test_exif_rotation_and_crop_geometry(self):
        path = self.root / "rotated.jpg"
        image = Image.new("RGB",(80,40),"white"); exif = image.getexif(); exif[274] = 6; image.save(path,exif=exif)
        report = extract.inspect_image(path,self.root,False)
        self.assertEqual(report.pages[0].metadata["exif_orientation"],6)
        with Image.open(report.pages[0].render_path) as image: self.assertEqual(image.size,(40,80))
        q=self.entry["details"]["questions"][0]
        cropped=annotations.crop_report(self.report,q)
        with Image.open(cropped.pages[0].render_path) as image: self.assertLess(image.height,200)

    def test_font_transparency_repeatability_and_missing_glyph(self):
        first=handwriting.render(self.root,mark="partial",seed="fixed")
        second=handwriting.render(self.root/"another",mark="partial",seed="fixed")
        self.assertEqual(first.read_bytes(),second.read_bytes())
        with Image.open(first) as im:
            self.assertEqual(im.mode,"RGBA")
            self.assertEqual(im.getchannel("A").getextrema(),(0,255))
        with self.assertRaises(local_workers.WorkerError): handwriting.render(self.root,text="\U0001f9ea")

    def test_export_preserves_ink_and_source_and_never_double_stamps(self):
        source=self.root/"source.pdf"; c=canvas.Canvas(str(source),pagesize=(400,500));c.drawString(50,370,"1. 2+3=5");c.save()
        writer=PdfWriter();writer.add_page(PdfReader(source).pages[0])
        ink=DictionaryObject({NameObject("/Type"):NameObject("/Annot"),NameObject("/Subtype"):NameObject("/Ink"),
             NameObject("/Rect"):RectangleObject([50,260,180,330]),NameObject("/InkList"):ArrayObject([ArrayObject([FloatObject(v) for v in [60,270,90,320,170,280]])]),
             NameObject("/C"):ArrayObject([FloatObject(0),FloatObject(0),FloatObject(1)])})
        writer.add_annotation(0,ink); writer.write(source)
        digest=hashlib.sha256(source.read_bytes()).hexdigest()
        report=extract.inspect_pdf(source,self.root/"_work",False)
        entry={**self.entry,"source_path":str(source),"pages":[asdict(p) for p in report.pages]}
        out=self.root/"marked.pdf"
        first=stamp.stamp_entry(entry,out,self.root/"_work")
        png=Path(first["marked_pngs"][0]).read_bytes()
        stamp.stamp_entry(entry,out,self.root/"_work")
        self.assertEqual(png,Path(first["marked_pngs"][0]).read_bytes())
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(),digest)
        result=PdfReader(out)
        self.assertEqual(len(result.pages),2)
        self.assertEqual(str(result.pages[0]["/Annots"][0].get_object()["/Subtype"]),"/Ink")

    def test_cropbox_all_pdf_rotations_place_inside_normalized_box(self):
        for rotation in (0,90,180,270):
            source=self.root/f"source-{rotation}.pdf"
            c=canvas.Canvas(str(source),pagesize=(400,500)); c.drawString(80,400,"Test"); c.save()
            pg=PdfReader(source).pages[0];pg.cropbox=RectangleObject([40,50,360,450]);pg.rotate(rotation)
            writer=PdfWriter();writer.add_page(pg);writer.write(source)
            report=extract.inspect_pdf(source,self.root/"_work",False)
            entry=copy.deepcopy(self.entry);entry.update(source_path=str(source),pages=[asdict(p) for p in report.pages])
            q=entry["details"]["questions"][0]
            q["placement"]={"page":1,"bbox":[.7,.65,.8,.75],"manual":True}
            result=stamp.stamp_entry(entry,self.root/f"out-{rotation}.pdf",self.root/"_work")
            self.assertEqual(len(result["placements"]),1)
            with Image.open(result["marked_pngs"][0]).convert("RGB") as image:
                pixels=image.load(); positions=[(x/image.width,y/image.height) for y in range(image.height) for x in range(image.width) if pixels[x,y][0]>120 and pixels[x,y][1]<100 and pixels[x,y][2]<100]
            self.assertTrue(positions)
            self.assertTrue(all(.697<=x<=.803 and .647<=y<=.753 for x,y in positions),rotation)

    def test_multiple_photos_and_cross_page_question_mark_once_at_end(self):
        other=self.root/"second.png";Image.new("RGB",(800,1000),"white").save(other)
        merged=extract.merge_reports([self.report,extract.inspect_image(other,self.root/"_work",False)],self.root)
        entry=copy.deepcopy(self.entry);entry.update(source_path=str(self.root),pages=[asdict(p) for p in merged.pages])
        q=entry["details"]["questions"][0];q["regions"].append({"page":2,"role":"answer","bbox":[.1,.1,.5,.3]})
        result=stamp.stamp_entry(entry,self.root/"multi.pdf",self.root/"_work")
        self.assertEqual(len(result["marked_pngs"]),3)
        self.assertEqual([p["page"] for p in result["placements"]],[2])

    def test_location_edit_never_scores_and_rejects_overlap(self):
        q=self.entry["details"]["questions"][0]
        from graders import providers
        with patch.object(providers,"complete",side_effect=AssertionError("must not grade")):
            edited=review.edit_question(self.entry,"q1",dict(action="location",placement={"page":1,"bbox":[.75,.3,.83,.37]}))
        self.assertEqual(edited["score"],3)
        with self.assertRaises(ValueError): review.edit_question(self.entry,"q1",dict(action="location",placement={"page":1,"bbox":[.2,.23,.3,.28]}))
        self.assertNotIn("placement",q)

    def test_excel_question_edit_and_legacy_total(self):
        batch=self.root/"batch.json";batch.write_text(json.dumps({"reports":[self.entry]}),encoding="utf-8")
        path=self.root/"edit.xlsx";summary.build_workbook([self.entry],path)
        from openpyxl import load_workbook
        wb=load_workbook(path);wb["逐题明细"].cell(2,5).value=1.5;wb.save(path);wb.close()
        updated=summary.merge_edits_into_batch(summary.read_workbook(path),batch)[0]
        self.assertEqual(updated["score"],1.5)
        self.assertEqual(updated["details"]["questions"][0]["mark"],"partial")
        updated=summary.merge_edits_into_batch([dict(source_file="student.png",score=0)],batch)[0]
        self.assertEqual(updated["score"],0)
        self.assertEqual(updated["details"]["questions"][0]["mark"],"correct")
        self.assertTrue(updated["manual_total_adjustment"])

    def test_excel_duplicate_filenames_need_ids(self):
        other={**self.entry,"report_id":"student2"}
        batch=self.root/"batch.json";batch.write_text(json.dumps({"reports":[self.entry,other]}),encoding="utf-8")
        with self.assertRaises(ValueError):summary.merge_edits_into_batch([dict(source_file="student.png",score=1)],batch)
        updated=summary.merge_edits_into_batch([dict(report_id="student2",score=1)],batch)
        self.assertEqual([r["score"] for r in updated],[3,1])

    def test_real_image_http_upload_review_preview_and_disk_restore(self):
        import app as web
        from graders import providers
        seg={"doc_type":"exam","questions":[
             {"no":"1","type":"fill","question":"2+3=?","student_answer":"5","full_score":3,"pages":[1],"answer_state":"present","transcription_confirmed":True,"location_confident":True,"regions":[{"page":1,"role":"answer","bbox":[.08,.22,.6,.31]}]},
             {"no":"2","type":"fill","question":"1+1=?","student_answer":"","full_score":2,"pages":[1],"answer_state":"unreadable","location_confident":True,"regions":[{"page":1,"role":"answer","bbox":[.08,.5,.6,.6]}]}]}
        client=web.app.test_client()
        with patch.object(config,"OUTPUT_DIR",self.root),patch.object(providers,"complete",side_effect=[json.dumps(seg),json.dumps({"answers":[{"no":"1","answer":"5"}]})]) as call:
            response=client.post("/grade",data={"task_type":"auto","provider":"deepseek","model":"deepseek-flash","api_key":"test-key","ocr_enabled":"false","reports":(io.BytesIO(self.path.read_bytes()),"student.png")},content_type="multipart/form-data")
            events=[json.loads(line[6:]) for line in response.data.decode().splitlines() if line.startswith("data: ")]
            done=next(e for e in events if e["type"]=="done")
            self.assertTrue(call.call_args_list[0].args[3])
            self.assertTrue(all(Path(p).is_file() for p in call.call_args_list[0].args[3]))
        job=done["job_id"];item=done["results"][0];rid=item["report_id"]
        self.assertTrue(item["score_provisional"])
        self.assertTrue(item["marked_pdf_path"],item.get("stamp_error"))
        self.assertEqual(item["details"]["questions"][0]["mark"],"correct")
        self.assertIsNone(item["details"]["questions"][1]["mark"])
        self.assertEqual(client.get(f"/preview/{job}/{rid}/1").status_code,200)
        self.assertEqual(client.get(f"/preview/{job}/{rid}/999").status_code,404)
        first_q=item["details"]["questions"][0]["id"];pending=item["details"]["questions"][1]["id"]
        with patch.object(providers,"complete",return_value=json.dumps({"answers":[{"no":"2","answer":"2"}]})) as call:
            response=client.post(f"/review/{job}/{rid}/{pending}",json={"action":"transcript","student_answer":"2","answer_state":"present","api_key":"test-key"})
            self.assertEqual(response.status_code,200,response.json)
            self.assertEqual(call.call_count,1)
        item=response.json["results"][0]
        self.assertEqual(item["score"],5)
        self.assertFalse(item["score_provisional"])
        with patch.object(providers,"complete",side_effect=AssertionError("manual edits must not call model")):
            response=client.post(f"/review/{job}/{rid}/{first_q}",json={"action":"score","score":1.5})
            self.assertEqual(response.status_code,200,response.json)
        self.assertEqual(response.json["results"][0]["details"]["questions"][0]["mark"],"partial")
        saved=Path(web.JOBS[job]["job_dir"])/"job.json"
        self.assertNotIn("test-key",saved.read_text(encoding="utf-8"))
        web.JOBS.pop(job)
        with patch.object(config,"OUTPUT_DIR",self.root):web.load_jobs_from_disk()
        self.assertEqual(client.get(f"/review/{job}/{rid}").status_code,200)
        web.JOBS.pop(job)


if __name__ == "__main__": unittest.main()
