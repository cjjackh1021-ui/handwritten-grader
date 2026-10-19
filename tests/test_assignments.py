"""作业题型(简答/数学/作文)离线测试: 用桩替换模型调用, 不需要 API Key。

运行: python tests/test_assignments.py   (也兼容 pytest)
"""

from __future__ import annotations

import io
import json
import shutil
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from graders import assignments, extract, math_check, providers  # noqa: E402
from graders.assignments import TaskSpec  # noqa: E402

CHOICE = providers.ModelChoice(provider="deepseek", api_key="sk-test", use_vision=False)


class Stub:
    """按顺序返回预设回复, 记录每次收到的 user prompt。"""

    def __init__(self, *replies):
        self.replies = list(replies)
        self.prompts: list[str] = []
        self.images: list = []

    def __call__(self, choice, system, user, images=None, max_tokens=None):
        self.prompts.append(user)
        self.images.append(images)
        reply = self.replies[min(len(self.prompts), len(self.replies)) - 1]
        return reply if isinstance(reply, str) else json.dumps(reply, ensure_ascii=False)


def with_stub(stub):
    original = providers.complete
    providers.complete = stub
    return original


def text_report(text: str, name: str = "张三") -> extract.ReportInfo:
    tmp = Path(tempfile.mkdtemp()) / f"2023001{name}.txt"
    tmp.write_text(text, encoding="utf-8")
    return extract.inspect_text_file(tmp)


def run(spec, report, *replies):
    stub = Stub(*replies)
    original = with_stub(stub)
    try:
        return assignments.grade_submission(CHOICE, spec, report), stub
    finally:
        providers.complete = original


# ---------------------------------------------------------------------------
def test_math_check():
    c = math_check.check_answer
    assert c("x=3", "3")["equal"]
    assert c("1/2", "0.5")["equal"]
    assert c("2√3", "sqrt(12)")["equal"]
    assert c("x=1或x=2", "2,1")["equal"]
    assert c("x=1", "x=2")["equal"] is False
    assert c("12.56平方厘米", "12.56")["equal"]
    assert c("增大", "增大")["equal"]
    assert c("", "3")["comparable"] is False
    assert c("在R上是增函数", "单调递增")["method"] == "text"  # 文字结论不当符号比较
    assert c("increasing", "monotonically increasing")["method"] == "text"
    poly = c("(x+1)^2", "x²+2x+1")  # 有 sympy 判等价; 没装时退化为文本比对(不作为推翻模型的依据)
    assert poly["equal"] is True or poly["method"] == "text"
    assert c("12.56", "4π")["approx"]  # 学生保留两位小数


def test_parse_points():
    pts = assignments.parse_points("1. 概念 | 2\n(2) 原理（4分）\n③ 实例：4分\n补充说明")
    assert [p["full"] for p in pts] == [2, 4, 4, None]
    scaled, note = assignments.scale_points(pts, 12)
    assert sum(p["full"] for p in scaled) == 12 and scaled[-1]["full"] == 2 and note
    scaled, note = assignments.scale_points([{"name": "a", "full": 3}, {"name": "b", "full": 3}], 10)
    assert [p["full"] for p in scaled] == [5, 5] and "缩放" in note
    spec = TaskSpec(task_type="essay", question="题", full_score=50)
    dims, _ = spec.points()
    assert len(dims) == 5 and sum(d["full"] for d in dims) == 50


SHORT = TaskSpec(task_type="short", question="简述光合作用的条件与产物", reference="需要光照和叶绿体, 产物为有机物和氧气",
                 full_score=10, points_text="需要光照 | 3\n需要叶绿体 | 3\n产物为有机物和氧气 | 4")


def short_reply(scores, total=None, comment="要点基本答全"):
    names = ["需要光照", "需要叶绿体", "产物为有机物和氧气"]
    fulls = [3, 3, 4]
    return {"points": [{"point": n, "full": f, "score": s, "hit": "full", "evidence": "原文"}
                       for n, f, s in zip(names, fulls, scores)],
            "score": sum(scores) if total is None else total, "comment": comment, "reason": "漏答叶绿体"}


def test_short_pass_first_try():
    r, stub = run(SHORT, text_report("光合作用需要光, 产生氧气和有机物"), short_reply([3, 0, 4]))
    assert r["score"] == 7 and r["full_score"] == 10 and r["steps"] == 1 and "warning" not in r
    assert [p["hit"] for p in r["details"]["points"]] == ["full", "miss", "full"]  # hit 由分数推导
    assert "需要叶绿体 —— 3分" in stub.prompts[0]


def test_short_retry_on_inconsistent_sum():
    r, stub = run(SHORT, text_report("光合作用需要光"),
                  short_reply([3, 0, 0], total=6),  # 分项和 3 ≠ 总分 6
                  short_reply([3, 0, 0]))
    assert r["steps"] == 2 and r["score"] == 3
    assert "不一致" in stub.prompts[1] and "上一次输出" in stub.prompts[1]
    assert any(t["type"] == "reflect" for t in r["trace"])


def test_short_repair_after_max_attempts():
    bad = short_reply([3, 3, 4], total=10, comment="这份作答非常完整而且条理清楚值得表扬继续保持下去")
    r, stub = run(SHORT, text_report("..."), bad)
    assert len(stub.prompts) == assignments.MAX_ATTEMPTS
    assert r["warning"] and len(r["comment"]) <= 20 and r["score"] == 10


MATH = TaskSpec(task_type="math", question="解方程 2x+1=5", reference="2x=4, x=2", full_score=6,
                points_text="移项 | 2\n系数化为1 | 2\n写出答案 | 2", final_answer="x=2")


def math_reply(student_final, final_correct, scores, corrects):
    return {"steps": [{"name": n, "full": 2, "score": s, "correct": c, "note": ""}
                      for n, s, c in zip(["移项", "系数化为1", "写出答案"], scores, corrects)],
            "first_error_step": 0, "student_final_answer": student_final, "reference_final_answer": "x=2",
            "final_correct": final_correct, "score": sum(scores), "comment": "计算出错", "reason": "第2步除法错误"}


def test_math_checker_overrules_model():
    report = text_report("2x=4\nx=3")
    r, stub = run(MATH, report,
                  math_reply("x=3", True, [2, 2, 2], [True, True, True]),     # 模型误判为正确且满分
                  math_reply("x=3", False, [2, 0, 1], [True, False, False]))
    assert r["steps"] == 2 and r["score"] == 3
    assert "不等价" in stub.prompts[1]
    d = r["details"]
    assert d["final_correct"] is False and d["first_error_step"] == 2
    assert d["answer_check"]["method"] == "numeric"
    assert any(t.get("name") == "check_final_answer" for t in r["trace"])


def test_math_equivalent_forms_accepted():
    r, _ = run(MATH, text_report("x = 4/2"), math_reply("x=4/2", True, [2, 2, 2], [True, True, True]))
    assert r["score"] == 6 and r["details"]["final_correct"] is True and r["steps"] == 1


def test_math_text_mismatch_not_authoritative():
    spec = TaskSpec(task_type="math", question="判断单调性", reference="单调递增", full_score=4,
                    points_text="求导 | 2\n下结论 | 2", final_answer="单调递增")
    reply = {"steps": [{"name": "求导", "full": 2, "score": 2, "correct": True},
                       {"name": "下结论", "full": 2, "score": 2, "correct": True}],
             "student_final_answer": "在R上是增函数", "final_correct": True, "score": 4,
             "comment": "正确", "reason": "过程完整"}
    r, stub = run(spec, text_report("f'(x)>0, 在R上是增函数"), reply)
    assert len(stub.prompts) == 1 and r["details"]["final_correct"] is True  # 文本不同不推翻模型


ESSAY = TaskSpec(task_type="essay", question="以「成长」为题写一篇记叙文", full_score=50, word_limit=600)


def essay_reply(norm_score):
    dims, _ = ESSAY.points()
    scores = {d["name"]: d["full"] * 0.8 for d in dims}
    norm = next(d for d in dims if "规范" in d["name"])
    scores[norm["name"]] = norm_score
    return {"off_topic": False,
            "dimensions": [{"name": n, "full": next(d["full"] for d in dims if d["name"] == n), "score": s,
                            "comment": "好"} for n, s in scores.items()],
            "highlights": ["开头设问引人入胜"], "suggestions": ["结尾再点题"],
            "revisions": [{"original": "我很开心。", "revised": "那一刻, 我心里像开了花。"}],
            "score": sum(scores.values()), "comment": "感情真挚", "reason": "内容充实, 字数不足"}


def test_essay_word_limit_enforced():
    dims, _ = ESSAY.points()
    norm_full = next(d["full"] for d in dims if "规范" in d["name"])
    r, stub = run(ESSAY, text_report("成长" * 100), essay_reply(norm_full), essay_reply(norm_full - 2))
    assert r["steps"] == 2 and "不应满分" in stub.prompts[1]
    d = r["details"]
    assert d["word_count"] == 200 and d["short_of_limit"] and d["language"] == "zh"
    assert d["revisions"][0]["revised"] and "约 200 字" in stub.prompts[0]


def test_essay_english_detected():
    spec = TaskSpec(task_type="essay", question="My hometown", full_score=25)
    dims, _ = spec.points()
    reply = {"dimensions": [{"name": d["name"], "full": d["full"], "score": d["full"]} for d in dims],
             "highlights": ["vivid"], "suggestions": ["more linking words"], "score": 25,
             "comment": "Great", "reason": "ok"}
    r, stub = run(spec, text_report("My hometown is a beautiful small city. " * 5), reply)
    assert r["details"]["language"] == "en" and r["details"]["word_count"] == 35
    assert "英语作文" in stub.prompts[0]


def test_invalid_json_retries_then_fails():
    stub_replies = ["抱歉我无法批改", "还是不行", "{bad json"]
    try:
        run(SHORT, text_report("x"), *stub_replies)
    except providers.ProviderError as exc:
        assert "JSON" in str(exc)
    else:
        raise AssertionError("应抛出 ProviderError")


def _png(path: Path, text: str):
    from PIL import Image, ImageDraw, ImageFont

    im = Image.new("RGB", (900, 200), "white")
    draw = ImageDraw.Draw(im)
    try:
        font = ImageFont.truetype("arial.ttf", 48)
    except OSError:
        font = ImageFont.load_default()
    draw.text((20, 60), text, fill="black", font=font)
    im.save(path)


def test_image_submission_and_merge():
    tmp = Path(tempfile.mkdtemp())
    a, b = tmp / "p1.png", tmp / "p2.png"
    _png(a, "Answer: x = 2")
    _png(b, "Check: 2*2+1 = 5")
    work = tmp / "_work"
    ia = extract.inspect_image(a, work, ocr_enabled=True)
    ib = extract.inspect_image(b, work, ocr_enabled=True)
    assert ia.pages[0].needs_visual and Path(ia.pages[0].render_path).is_file()
    folder = tmp / "2023000005李四"
    merged = extract.merge_reports([ia, ib], folder)
    assert merged.page_count == 2 and merged.student_name == "李四" and merged.student_id == "2023000005", \
        (merged.student_name, merged.student_id)
    assert [p.page_num for p in merged.pages] == [1, 2]

    # 无视觉 + 无 OCR 文本: 明确报错而不是瞎批
    blank = extract.inspect_image(a, work, ocr_enabled=False)
    try:
        run(MATH, blank, "{}")
    except providers.ProviderError as exc:
        assert "视觉" in str(exc)
    else:
        raise AssertionError("应提示需要视觉模型")

    # 有视觉: 图片随请求发送
    vision_choice = providers.ModelChoice(provider="openai", api_key="sk", use_vision=True)
    stub = Stub(math_reply("x=2", True, [2, 2, 2], [True, True, True]))
    original = with_stub(stub)
    try:
        r = assignments.grade_submission(vision_choice, MATH, merged)
    finally:
        providers.complete = original
    assert r["score"] == 6 and len(stub.images[0]) == 2 and "以图片内容为准" in stub.prompts[0]


def test_flask_end_to_end():
    import app as webapp
    from openpyxl import load_workbook

    stub = Stub(short_reply([3, 3, 2]))
    original = with_stub(stub)
    client = webapp.app.test_client()
    files = [
        (io.BytesIO("光照、叶绿体".encode()), "作业/2023000001张三.txt"),
        (io.BytesIO("需要光".encode()), "作业/2023000002李四/第1页.txt"),
        (io.BytesIO("产物是氧气".encode()), "作业/2023000002李四/第2页.txt"),
    ]
    data = {"task_type": "short", "question": SHORT.question, "reference": SHORT.reference,
            "points": SHORT.points_text, "full_score": "10", "provider": "deepseek", "api_key": "sk-test",
            "use_vision": "false", "ocr_enabled": "false", "grading_criteria": "",
            "reports": [(f, name) for f, name in files]}
    try:
        resp = client.post("/grade", data=data, content_type="multipart/form-data")
        assert resp.status_code == 200, resp.get_data(as_text=True)
        events = [json.loads(line[5:]) for line in resp.get_data(as_text=True).split("\n\n") if line.startswith("data:")]
    finally:
        providers.complete = original
    done = next(e for e in events if e["type"] == "done")
    results = done["results"]
    assert len(results) == 2, [r["source_file"] for r in results]  # 李四两页合并为一份
    by_name = {r["student_name"]: r for r in results}
    assert set(by_name) == {"张三", "李四"}, [(r["student_name"], r["student_id"], r["source_file"]) for r in results]
    assert by_name["李四"]["student_id"] == "2023000002"
    assert all(r["task_type"] == "short" and r["full_score"] == 10 and r["score"] == 8 for r in results)
    assert done["statistics"]["full_score"] == 10
    lisi_prompt = next(p for p in stub.prompts if "李四" in p)
    assert "需要光" in lisi_prompt and "产物是氧气" in lisi_prompt

    job = webapp.JOBS[done["job_id"]]
    ws = load_workbook(job["xlsx"]).active
    header = [c.value for c in ws[1]]
    row = [c.value for c in ws[2]]
    assert header[2:5] == ["分数", "满分", "题型"] and row[3] == 10 and row[4] == "简答题"
    assert job["options"]["task"]["task_type"] == "short"

    bad = client.post("/grade", data={"task_type": "math", "question": "题", "api_key": "sk"},
                      content_type="multipart/form-data")
    assert bad.status_code == 400 and "参考解答" in bad.get_json()["error"]
    shutil.rmtree(job["job_dir"], ignore_errors=True)


def test_lab_path_unchanged():
    from graders import grading

    reply = {"score": 88, "comment": "好", "reason": "完整", "dimensions": [{"name": "原理", "score": 18, "full": 20}]}
    stub = Stub(reply)
    original = with_stub(stub)
    try:
        r = grading.grade_report(CHOICE, "原理 20分", text_report("实验报告内容"))
    finally:
        providers.complete = original
    assert r["score"] == 88 and r["dimensions"][0]["full"] == 20


if __name__ == "__main__":
    tests = [(n, f) for n, f in sorted(globals().items()) if n.startswith("test_") and callable(f)]
    failed = 0
    for name, fn in tests:
        try:
            fn()
            print(f"PASS  {name}")
        except Exception as exc:  # noqa: BLE001
            failed += 1
            import traceback

            print(f"FAIL  {name}: {exc!r}")
            traceback.print_exc()
    print(f"\n{len(tests) - failed}/{len(tests)} passed")
    sys.exit(1 if failed else 0)
