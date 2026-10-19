"""Behavioral coverage for the installed skill's API-free workflow."""
import copy
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from PIL import Image, ImageDraw
from openpyxl import load_workbook
from pypdf import PdfReader
from graders import annotations, codex_handwriting as hw, providers, summary


class CodexHandwritingTests(unittest.TestCase):
    def setUp(self):
        temp = tempfile.TemporaryDirectory(prefix="codex-handwriting-")
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.files = []
        for n in (1, 2):
            path = self.root / f"原稿{n}.png"
            image = Image.new("RGB", (800, 1100), "white")
            ImageDraw.Draw(image).text((80, 220), f"SYNTHETIC PAGE {n} - 2+3=5", fill="black")
            image.save(path)
            self.files.append(path)
        manifest = self.root / "input.json"
        hw.write(manifest, {"submissions": [{"student_name": "测试", "student_id": "001", "files": list(map(str, self.files))}]})
        prepared = hw.prepare(manifest, self.root / "prepare", "off")
        self.packet = Path(prepared["packet"])
        self.rid = hw.read(self.packet)["reports"][0]["report_id"]
        rows = []
        for n, state in ((1, "present"), (2, "unreadable")):
            rows.append(dict(no=str(n), type="math", question="计算 2+3", full=3, student_answer="5" if n == 1 else "?",
                             answer_state=state, transcription_confirmed=n == 1, location_confident=True,
                             regions=[dict(page=n, role="answer", bbox=[.08,.18,.6,.28])]))
        questions = self.root / "questions.json"
        hw.write(questions, {"reports": [dict(report_id=self.rid, questions=rows)]})
        result = hw.locate(self.packet, questions, self.root / "layout")
        self.reports = hw.read(result["layout"])["reports"]
        self.grades = hw.read(result["grading_template"])
        self.grades["grades"][0].update(score=3, grading_confirmed=True, reason="原图结果为 5，正确")

    def test_real_pipeline_no_api_and_draft_exports(self):
        with patch.object(providers, "chat", side_effect=AssertionError("No external API allowed")):
            scored = hw.apply_grades(self.reports, self.grades)
            result = hw.export(scored, self.root / "export")
        self.assertEqual(result["marked"], 1)
        self.assertEqual(result["failures"], [])
        exported = hw.read(result["batch_results_path"])["reports"][0]
        qs = exported["details"]["questions"]
        self.assertEqual(qs[0]["mark"], "correct")
        self.assertIsNone(qs[1]["mark"])
        self.assertTrue(exported["score_provisional"])
        self.assertEqual(exported["score"], 3)
        self.assertEqual(len(PdfReader(exported["marked_pdf_path"]).pages), 3)
        self.assertTrue(all(Path(p).exists() for p in exported["marked_pngs"]))
        with zipfile.ZipFile(result["zip_path"]) as z:
            self.assertIn("批改结果.xlsx", z.namelist())
            self.assertIn("review_pending.json", z.namelist())
            self.assertEqual(len([n for n in z.namelist() if n.endswith(".pdf")]), 1)
        workbook = load_workbook(result["excel"])
        self.assertEqual(workbook["逐题明细"].max_row, 3)
        workbook.close()

    def test_only_corrected_question_regraded_with_new_digest(self):
        scored = hw.apply_grades(self.reports, self.grades)
        first = copy.deepcopy(scored[0]["details"]["questions"][0])
        qid = scored[0]["details"]["questions"][1]["id"]
        changed = hw.revise(scored, {"changes": [dict(report_id=self.rid, question_id=qid, action="transcript", student_answer="4", answer_state="present")]})
        self.assertEqual(changed[0]["details"]["questions"][0], first)
        self.assertTrue(changed[0]["details"]["questions"][1]["grading_review"])
        stale = {"grades": [self.grades["grades"][1]]}
        with self.assertRaisesRegex(ValueError, "已变化"):
            hw.apply_grades(changed, stale, False)
        grades = hw.grade_template(changed, {(self.rid, qid)})
        grades["grades"][0].update(score=0, grading_confirmed=True, reason="结果错误")
        final = hw.apply_grades(changed, grades, False)
        self.assertEqual(final[0]["details"]["questions"][1]["mark"], "incorrect")
        self.assertEqual(final[0]["details"]["questions"][0], first)

    def test_uncertain_transcription_cannot_be_scored_as_zero_cross(self):
        self.grades["grades"][1].update(score=0, grading_confirmed=True, reason="Wrong")
        result = hw.apply_grades(self.reports, self.grades)
        self.assertIsNone(result[0]["details"]["questions"][1]["mark"])

    def test_reject_missing_duplicate_stale_or_out_of_range_grades(self):
        variants = [dict(grades=self.grades["grades"][:1]), dict(grades=[self.grades["grades"][0]] * 2)]
        for field, value in (("score", 4), ("score", float("nan")), ("question_digest", "old")):
            data = copy.deepcopy(self.grades)
            data["grades"][0][field] = value
            variants.append(data)
        for data in variants:
            with self.assertRaises(ValueError):
                hw.apply_grades(self.reports, data)

    def test_source_changes_prevent_export_and_existing_output_is_kept(self):
        self.files[0].write_bytes(b"changed")
        with self.assertRaisesRegex(ValueError, "原稿"):
            hw.export(self.reports, self.root / "out")
        marker = self.root / "teacher.xlsx"
        marker.write_bytes(b"teacher edits")
        with self.assertRaises(ValueError):
            hw.fresh_output(self.root)
        self.assertEqual(marker.read_bytes(), b"teacher edits")

    def test_multi_page_and_cross_page_one_mark(self):
        r = copy.deepcopy(self.reports[0])
        q = r["details"]["questions"][0]
        q["regions"].append(dict(page=2, role="answer", bbox=[.08,.18,.6,.28]))
        q.update(score=1.5, grading_review=False, recognition_review=False, position_review=False)
        r["details"]["questions"] = [q]
        result = hw.export([r], self.root / "crosspage")
        exported = hw.read(result["batch_results_path"])["reports"][0]
        self.assertEqual(len(exported["placements"]), 1)
        self.assertEqual(exported["placements"][0]["page"], 2)
        self.assertEqual(exported["placements"][0]["mark"], "partial")

    def test_excel_total_override_survives_regeneration(self):
        scored = hw.apply_grades(self.reports, self.grades)
        first = hw.export(scored, self.root / "export")
        workbook = load_workbook(first["excel"])
        workbook["批改结果"].cell(2, 3, 4)
        edited = self.root / "edited.xlsx"
        workbook.save(edited)
        workbook.close()
        merged = summary.merge_edits_into_batch(summary.read_workbook(edited), Path(first["batch_results_path"]))
        regenerated = hw.export(merged, self.root / "regenerated")
        entry = hw.read(regenerated["batch_results_path"])["reports"][0]
        self.assertEqual(entry["score"], 4)
        self.assertEqual(entry["details"]["questions"][0]["score"], 3)
        self.assertEqual(len(PdfReader(entry["marked_pdf_path"]).pages), 3)


if __name__ == "__main__":
    unittest.main()
